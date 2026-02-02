from docx import Document
from docx.shared import Pt
from datetime import datetime
from copy import deepcopy


# -----------------------------
# PERIOD FORMATTER
# -----------------------------
def format_period_from_dates(dates):
    if not dates:
        return ""

    dates = sorted(dates)
    first = dates[0]
    last = dates[-1]

    month = first.strftime("%B")
    year = first.year

    if first.day <= 15 and last.day <= 15:
        return f"For the Period: {month} 1-15, {year}"
    elif first.day >= 16:
        return f"For the Period: {month} 16-31, {year}"
    else:
        return f"For the Period: {month} {year}"


# -----------------------------
# HELPER: INSERT ROW
# -----------------------------
def insert_row_after(table, index):
    """
    Inserts a new row by cloning the layout of the row at 'index'.
    This preserves column widths and merged cells.
    """
    row_to_clone = table.rows[index]
    
    # Deep copy the XML element of the row to preserve layout (cols/merges)
    new_tr = deepcopy(row_to_clone._tr)
    
    # Insert the new row XML after the current row
    row_to_clone._tr.addnext(new_tr)
    
    # Access the newly created row (it will be at index + 1)
    new_row = table.rows[index + 1]
    
    # Clear content in the cloned row
    for cell in new_row.cells:
        cell.text = ""
        
    return new_row


# -----------------------------
# BUILD TASK LIST
# -----------------------------
def build_tasks_from_attendance(employee_data, overrides=None):
    overrides = overrides or {}
    override_tasks = overrides.get("tasks", {})

    month_name = employee_data.get("month_name")
    year = employee_data.get("year")

    if not month_name or not year:
        return []

    tasks = []

    for day_str, data in employee_data.items():
        if not day_str.isdigit():
            continue

        if not any([
            data.get("am_in"), data.get("am_out"),
            data.get("pm_in"), data.get("pm_out"),
        ]):
            continue

        try:
            dt = datetime.strptime(
                f"{int(day_str)} {month_name} {year}",
                "%d %B %Y"
            ).date()
        except Exception:
            continue

        tasks.append({
            "date": dt,
            "task": override_tasks.get(day_str, "")
        })

    tasks.sort(key=lambda x: x["date"])
    return tasks


# -----------------------------
# MAIN GENERATOR
# -----------------------------
def generate_ar_docx(
    template_path,
    output_path,
    employee_data,
    parsed_attendance,
    overrides=None
):
    overrides = overrides or {}
    doc = Document(template_path)
    
    if not doc.tables:
        raise Exception("Template has no tables.")
        
    table = doc.tables[0]

    # --- 1. PREPARE DATA ---
    name = overrides.get("name", employee_data.get("name", ""))
    position = overrides.get("position", employee_data.get("position", ""))
    office = overrides.get("office", employee_data.get("office", ""))
    project = overrides.get("project", employee_data.get("project", ""))
    submitted_by = overrides.get("submitted_by", name)
    approved_by = overrides.get("approved_by", "")
    
    # Default title if not provided
    approver_title = overrides.get("approver_title", "PROVINCIAL OFFICER, ISABELA - CAUAYAN II")

    # --- 2. FILL HEADER INFO ---
    def fill_label(label_text, value, is_position=False):
        val_upper = value.upper()
        for row in table.rows:
            for cell in row.cells:
                if label_text in cell.text.upper():
                    if val_upper in cell.text.upper():
                        return
                    
                    p = cell.paragraphs[0]
                    
                    if is_position:
                        # Position goes on a new line
                        run = p.add_run(f"\n{val_upper}")
                    else:
                        # Others go beside with a space
                        run = p.add_run(f"  {val_upper}")
                        
                    run.bold = True
                    run.font.size = Pt(11)
                    return

    # Updated: Name now uses is_position=True to put text on a newline
    fill_label("NAME:", name, is_position=True)
    fill_label("POSITION:", position, is_position=True)
    fill_label("OFFICE:", office)
    fill_label("PROJECT:", project)

    # --- 3. FILL PERIOD ---
    if "period_text" in overrides and overrides["period_text"]:
        period_text = overrides['period_text']
    else:
        tasks = build_tasks_from_attendance(parsed_attendance, overrides)
        dates = [t["date"] for t in tasks]
        period_text = format_period_from_dates(dates)

    period_filled = False

    # Check Document Body
    for p in doc.paragraphs:
        if "for the period" in p.text.lower():
            p.text = "For the period of "
            run = p.add_run(period_text)
            run.bold = True
            p.alignment = 1
            period_filled = True
            break
            
    # Check Inside Table (Header rows) if not found in body
    if not period_filled:
        for row in table.rows[:5]:
            for cell in row.cells:
                if "for the period" in cell.text.lower():
                    cell.text = "For the period of "
                    p = cell.paragraphs[0]
                    p.alignment = 1
                    run = p.add_run(period_text)
                    run.bold = True
                    period_filled = True
                    break
                elif "accomplishment report" in cell.text.lower():
                    if "for the period" not in cell.text.lower():
                        p = cell.add_paragraph()
                        p.alignment = 1
                        p.add_run("For the period of ")
                        run = p.add_run(period_text)
                        run.bold = True
                        period_filled = True
                        break
            if period_filled: break

    # --- 4. INSERT TASKS DYNAMICALLY ---
    tasks = build_tasks_from_attendance(parsed_attendance, overrides)
    
    task_header_idx = -1
    footer_start_idx = -1

    for i, row in enumerate(table.rows):
        cell_text = row.cells[0].text.upper()
        if "TASKS" in cell_text:
            task_header_idx = i
        elif "SUBMITTED BY" in cell_text or "APPROVED BY" in cell_text:
            footer_start_idx = i
            break 

    if task_header_idx == -1 or footer_start_idx == -1:
        print("Warning: Could not find TASKS or SUBMITTED BY markers. Appending to end.")
        current_row_idx = len(table.rows)
    else:
        current_row_idx = task_header_idx + 1

    for task in tasks:
        date_str = task["date"].strftime("%B %d, %Y")
        task_desc = task["task"]

        if current_row_idx >= footer_start_idx:
            new_row = insert_row_after(table, current_row_idx - 1)
            row = new_row
            footer_start_idx += 1 
        else:
            row = table.rows[current_row_idx]

        row.cells[0].text = date_str
        row.cells[1].text = task_desc
        
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = 0 
                for run in p.runs:
                    run.font.size = Pt(10)

        current_row_idx += 1

    # --- 5. FILL SIGNATORIES (Smart Column Detection) ---
    if footer_start_idx != -1:
        sig_row = table.rows[footer_start_idx]
        
        sub_col_idx = -1
        app_col_idx = -1
        
        for idx, cell in enumerate(sig_row.cells):
            txt = cell.text.upper()
            if "SUBMITTED" in txt:
                sub_col_idx = idx
            elif "APPROVED" in txt:
                app_col_idx = idx
        
        if sub_col_idx == -1: sub_col_idx = 0
        if app_col_idx == -1 and len(sig_row.cells) > 1: app_col_idx = 1

        # Write Submitted By (Employee)
        if sub_col_idx < len(sig_row.cells):
            cell_sub = sig_row.cells[sub_col_idx]
            p_sub = cell_sub.add_paragraph()
            p_sub.alignment = 1 
            # 2 newlines before for signature space
            run_sub = p_sub.add_run(f"\n\n{submitted_by.upper()}\n\n")
            run_sub.bold = True
            run_sub.font.size = Pt(11)
        
        # Write Approved By (Provincial Officer)
        if app_col_idx != -1 and app_col_idx < len(sig_row.cells) and approved_by:
            cell_app = sig_row.cells[app_col_idx]
            p_app = cell_app.add_paragraph()
            p_app.alignment = 1 
            
            # Name
            run_app = p_app.add_run(f"\n\n{approved_by.upper()}")
            run_app.bold = True
            run_app.font.size = Pt(11)
            
            # Title (New line under name)
            # Updated: Added \n\n after title to add space at the bottom
            run_title = p_app.add_run(f"\n{approver_title}\n\n")
            run_title.bold = True 
            run_title.font.size = Pt(11)

    doc.save(output_path)
    return output_path