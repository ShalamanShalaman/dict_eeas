import re
import calendar
from io import BytesIO
from flask import current_app
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, Border, Side
from datetime import datetime, timedelta

from ar_utils import generate_ar_docx

def to_24h(time_str, is_pm, is_am):
    try:
        dt = datetime.strptime(time_str, "%H:%M:%S")
    except ValueError:
        try:
            dt = datetime.strptime(time_str, "%H:%M")
        except ValueError:
            return time_str 

    hour = dt.hour
    
    if is_pm and hour < 12:
        hour += 12
    elif is_am and hour == 12:
        hour = 0
        
    return dt.replace(hour=hour).strftime("%H:%M:%S")

def clean_daily_logs(logs):
    if not logs:
        return []

    def parse_time(t_str):
        try:
            return datetime.strptime(t_str, "%H:%M:%S")
        except ValueError:
            return datetime.strptime(t_str, "%H:%M")

    logs.sort(key=lambda x: parse_time(x[1]))

    cleaned = []
    last_time = None

    for action, time_str in logs:
        current_dt = parse_time(time_str)
        
        if last_time:
            diff = (current_dt - last_time).total_seconds()
            if diff < 60:
                continue

        cleaned.append((action, time_str))
        last_time = current_dt

    return cleaned

def parse_employees_data(text):
    employees = {}
    lines = text.split('\n')

    current_employee = None
    employee_data = {}
    day_checkins = {} 

    current_month_name = ""
    current_year = ""

    def process_employee_data(emp_name, emp_data, day_checks, month_name, year):
        if not emp_name or not day_checks:
            return

        if month_name:
            emp_data['month_name'] = month_name
        if year:
            emp_data['year'] = year

        for day, raw_logs in day_checks.items():
            if day not in emp_data:
                continue

            cleaned_logs = clean_daily_logs(raw_logs)

            for idx, (action, time) in enumerate(cleaned_logs):
                display_time = ':'.join(time.split(':')[:2])

                if action == 'C/IN':
                    if idx == 0: 
                        emp_data[day]['am_in'] = display_time
                    else: 
                        emp_data[day]['pm_in'] = display_time

                elif action == 'C/OUT':
                    if idx == 1 or (idx > 0 and not emp_data[day]['am_out']):
                        emp_data[day]['am_out'] = display_time
                    else:
                        emp_data[day]['pm_out'] = display_time

    for line in lines:
        line = line.strip()
        if not line:
            continue

        try:
            match = re.match(r'^([A-Za-z0-9\s,]+)\((\d+)\)$', line)
            if match:
                process_employee_data(
                    current_employee,
                    employee_data,
                    day_checkins,
                    current_month_name,
                    current_year
                )

                if current_employee and employee_data:
                    employees[current_employee] = employee_data

                current_employee = match.group(1).strip()
                employee_data = {}
                day_checkins = {}
                current_month_name = ""
                current_year = ""
                continue

            if not current_employee:
                continue

            parts = line.split()
            if len(parts) < 3:
                continue

            if '/' in parts[0] and parts[0].count('/') == 2:
                date_str = parts[0]
                time_part = parts[1] 
                
                rest_of_line = " ".join(parts[2:]).lower()
                
                is_pm = 'pm' in rest_of_line
                is_am = 'am' in rest_of_line
                
                action = 'UNKNOWN'
                if 'c/in' in rest_of_line:
                    action = 'C/IN'
                elif 'c/out' in rest_of_line:
                    action = 'C/OUT'
                
                time_part_24 = to_24h(time_part, is_pm, is_am)

                day_str, month_str, year_str = date_str.split('/')

                if not current_month_name:
                    current_month_name = calendar.month_name[int(month_str)]
                if not current_year:
                    current_year = year_str

                day_num = int(day_str)
                day_key = str(day_num)

                if day_key not in employee_data:
                    employee_data[day_key] = {
                        'am_in': '', 'am_out': '', 'pm_in': '', 'pm_out': '',
                        'undertime_hrs': '', 'undertime_min': '', 'remarks': ''
                    }
                    day_checkins[day_key] = []

                if action != 'UNKNOWN':
                    day_checkins[day_key].append((action, time_part_24))

        except Exception:
            continue

    process_employee_data(
        current_employee,
        employee_data,
        day_checkins,
        current_month_name,
        current_year
    )

    if current_employee and employee_data:
        employees[current_employee] = employee_data

    return employees

def generate_dtr(employee_name, employee_data, template_path=None, approver_name="", period_text=""):

    if template_path is None:
        template_path = current_app.config.get('DTR_TEMPLATE')

    if not template_path:
        raise FileNotFoundError("DTR template path is not configured")

    raw_data = dict(employee_data)
    month_name = raw_data.pop('month_name', '')
    year = raw_data.pop('year', '')

    employee_data = {
        int(k): v for k, v in raw_data.items() if k.isdigit()
    }

    wb = load_workbook(template_path)
    ws = wb.active

    for cell in ['C6', 'K6', 'C55', 'K55']:
        ws[cell] = employee_name

    header_val = period_text if period_text else (f"{month_name} {year}" if month_name else "")
    
    if header_val:
        for cell in ['E8', 'M8']:
            ws[cell] = header_val
            ws[cell].alignment = Alignment(horizontal='center', vertical='center')

    if approver_name:
        for cell in ['C61', 'K61']:
            ws[cell] = approver_name
            ws[cell].alignment = Alignment(horizontal='center', vertical='bottom')

    center = Alignment(horizontal='center', vertical='center')
    bold_font = Font(bold=True)
    
    remarks_map = {d: employee_data.get(d, {}).get('remarks', '').strip() for d in range(1, 32)}
    processed_remarks_days = set()

    for day in range(1, 32):
        row = 13 + day
        if row > 44: break
        
        day_data = employee_data.get(day, {})
        current_remark = remarks_map.get(day, "")

        ws[f'B{row}'] = day
        ws[f'J{row}'] = day

        if day in processed_remarks_days:
            continue

        if current_remark:
            span = 1
            for lookahead in range(day + 1, 32):
                if remarks_map.get(lookahead) == current_remark:
                    span += 1
                else:
                    break
            
            for d in range(day, day + span):
                processed_remarks_days.add(d)
            
            end_row = row + span - 1
            
            ws.merge_cells(start_row=row, start_column=3, end_row=end_row, end_column=6)
            cell_left = ws[f'C{row}']
            cell_left.value = current_remark
            cell_left.alignment = center
            cell_left.font = bold_font 

            ws.merge_cells(start_row=row, start_column=11, end_row=end_row, end_column=14)
            cell_right = ws[f'K{row}']
            cell_right.value = current_remark
            cell_right.alignment = center
            cell_right.font = bold_font

            for r in range(row, end_row + 1):
                ws[f'G{r}'] = ""
                ws[f'H{r}'] = ""
                ws[f'O{r}'] = ""
                ws[f'P{r}'] = ""

        else:
            updates = {
                'C': day_data.get('am_in', ''),
                'D': day_data.get('am_out', ''),
                'E': day_data.get('pm_in', ''),
                'F': day_data.get('pm_out', ''),
                'G': day_data.get('undertime_hrs', ''),
                'H': day_data.get('undertime_min', ''),
                
                'K': day_data.get('am_in', ''),
                'L': day_data.get('am_out', ''),
                'M': day_data.get('pm_in', ''),
                'N': day_data.get('pm_out', ''),
                'O': day_data.get('undertime_hrs', ''),
                'P': day_data.get('undertime_min', '')
            }

            for col, value in updates.items():
                cell = ws[f'{col}{row}']
                cell.value = value
                cell.alignment = center
                cell.font = bold_font

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generate_ar_from_employee_data(
    employee_name,
    employee_data,
    employee_position,
    employee_office,
    project,
    output_path,
    overrides=None,
    period_text=""
):

    template_path = current_app.config.get("AR_TEMPLATE")

    if not template_path:
        raise FileNotFoundError("AR template not configured")

    employee_info = {
        "name": employee_name,
        "position": employee_position,
        "office": employee_office,
        "project": project
    }
    
    current_overrides = overrides or {}
    if period_text:
        current_overrides["period_text"] = period_text

    return generate_ar_docx(
        template_path=template_path,
        output_path=output_path,
        employee_data=employee_info,
        parsed_attendance=employee_data,
        overrides=current_overrides
    )