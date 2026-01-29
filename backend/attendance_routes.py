from flask import Blueprint, request, jsonify, render_template, send_file, current_app
import PyPDF2
import os
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from attendance_utils import parse_employees_data

# Define the Blueprint
attendance_bp = Blueprint('attendance', __name__)

@attendance_bp.route('/', methods=['GET', 'POST'])
def index():
    extracted_text = ""
    employees = {}
    counts = {}
    error_message = ""

    if request.method == 'POST':
        if 'pdf_file' not in request.files:
            error_message = "No file part"
        else:
            file = request.files['pdf_file']
            if file.filename == '':
                error_message = "No selected file"
            else:
                # Use current_app.config to access the upload folder defined in app.py
                filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], file.filename)
                file.save(filepath)

                try:
                    pdf_file = open(filepath, 'rb')
                    reader = PyPDF2.PdfReader(pdf_file)
                    for page in reader.pages:
                        extracted_text += page.extract_text() + "\n"
                    pdf_file.close()

                    employees = parse_employees_data(extracted_text)

                    if not employees:
                        error_message = "No valid employee attendance data found in the PDF. Please ensure the PDF contains attendance records in the expected format."
                    else:
                        for name, data in employees.items():
                            day_count = sum(1 for k in data.keys() if k.isdigit())
                            counts[name] = day_count
                except Exception as e:
                    error_message = f"Error processing PDF: {str(e)}"

    return render_template('index.html', extracted_text=extracted_text, employees=employees, counts=counts, error_message=error_message)


@attendance_bp.route('/api/upload-attendance', methods=['POST'])
def api_upload_attendance():
    if 'attendanceFile' not in request.files:
        return jsonify({"error": "No file part"}), 400
    
    file = request.files['attendanceFile']
    if file.filename == '':
        return jsonify({"error": "No selected file"}), 400

    try:
        # Use current_app.config to access the upload folder defined in app.py
        filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], file.filename)
        file.save(filepath)

        extracted_text = ""
        with open(filepath, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                extracted_text += page.extract_text() + "\n"
        
        # Reuse the parsing function imported from attendance_utils.py
        employees_data = parse_employees_data(extracted_text)

        if not employees_data:
            return jsonify({"error": "No valid data found in PDF"}), 422

        return jsonify({
            "message": "Upload successful",
            "data": employees_data
        }), 200

    except Exception as e:
        return jsonify({"error": str(e)}), 500
    

@attendance_bp.route('/api/download-dtr', methods=['POST'])
def download_dtr():
    data = request.json

    employee_name = data.get("employee_name", "")

    try:
        raw_data = data.get("employee_data", {})

        month_name = raw_data.pop('month_name', '')
        year = raw_data.pop('year', '')

        employee_data = {
            int(k): v for k, v in raw_data.items() if k.isdigit()
        }

    except Exception:
        employee_data = {}
        month_name = ""
        year = ""

    # Ensure this path is correct relative to where you run python app.py
    template_path = 'templates_excel/template.xlsx'
    
    try:
        wb = load_workbook(template_path)
    except FileNotFoundError:
        return jsonify({"error": "Excel template not found. Please check 'templates_excel/template.xlsx'"}), 500

    ws = wb.active

    ws['C6'] = employee_name
    ws['K6'] = employee_name
    ws['C55'] = employee_name
    ws['K55'] = employee_name

    
    if month_name:
        ws['E8'] = f"{month_name} {year}"
        ws['E8'].alignment = Alignment(horizontal='center', vertical='center')

        ws['M8'] = f"{month_name} {year}"
        ws['M8'].alignment = Alignment(horizontal='center', vertical='center')
    
    center_aligned = Alignment(horizontal='center', vertical='center')
    bold_font = Font(bold=True)
    
    for day, day_data in employee_data.items():
        row = 13 + day
        
        if row <= 44:
            cell_updates = {
                'B': day,
                'C': day_data.get('am_in', ''),
                'D': day_data.get('am_out', ''),
                'E': day_data.get('pm_in', ''),
                'F': day_data.get('pm_out', ''),
                'G': day_data.get('undertime_hrs', ''),
                'H': day_data.get('undertime_min', ''),

                'J': day,
                'K': day_data.get('am_in', ''),
                'L': day_data.get('am_out', ''),
                'M': day_data.get('pm_in', ''),
                'N': day_data.get('pm_out', ''),
                'O': day_data.get('undertime_hrs', ''),
                'P': day_data.get('undertime_min', '')
            }
            
            for col, value in cell_updates.items():
                cell = ws[f'{col}{row}']
                cell.value = value
                cell.alignment = center_aligned
                if col != 'B':
                    cell.font = bold_font
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    clean_name = employee_name.replace(',', '').replace(' ', '_')
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'DTR_{clean_name}.xlsx'
    )