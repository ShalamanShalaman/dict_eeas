import os
import re
import calendar
import copy
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from datetime import datetime
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

from ar_utils import generate_ar_docx

def to_24h(time_str, is_pm, is_am):
    try:
        dt = datetime.strptime(time_str, "%H:%M:%S")
    except ValueError:
        try:
            dt = datetime.strptime(time_str, "%H:%M")
        except ValueError:
            return time_str 

    hour = dt.hour
    
    if is_pm and hour < 12:
        hour += 12
    elif is_am and hour == 12:
        hour = 0
        
    return dt.replace(hour=hour).strftime("%H:%M:%S")

def to_12h(time_str):
    if not time_str:
        return ""
    try:
        dt = datetime.strptime(str(time_str).strip(), "%H:%M")
        return dt.strftime("%I:%M")
    except ValueError:
        try:
            dt = datetime.strptime(str(time_str).strip(), "%H:%M:%S")
            return dt.strftime("%I:%M")
        except ValueError:
            return time_str

def clean_daily_logs(logs):
    if not logs:
        return []

    def parse_time(t_str):
        try:
            return datetime.strptime(t_str, "%H:%M:%S")
        except ValueError:
            return datetime.strptime(t_str, "%H:%M")

    logs.sort(key=lambda x: parse_time(x[1]))

    cleaned = []
    last_time = None

    for action, time_str in logs:
        current_dt = parse_time(time_str)
        
        if last_time:
            diff = (current_dt - last_time).total_seconds()
            if diff < 60:
                continue

        cleaned.append((action, time_str))
        last_time = current_dt

    return cleaned

def parse_employees_data(text, date_format='DMY'):
    employees = {}
    lines = text.split('\n')

    current_employee = None
    employee_data = {} 
    day_checkins = {} 

    def process_employee_data(emp_name, emp_data, day_checks):
        if not emp_name or not day_checks:
            return

        for m_key, m_checks in day_checks.items():
            for day, raw_logs in m_checks.items():
                if day not in emp_data[m_key]:
                    continue
                
                cleaned_logs = clean_daily_logs(raw_logs)
                for action, time in cleaned_logs:
                    display_time = to_12h(time)
                    
                    try:
                        hour = int(time.split(':')[0])
                    except ValueError:
                        hour = 0
                        
                    if action == 'C/IN':
                        if hour < 12:
                            emp_data[m_key][day]['am_in'].append(display_time)
                        else:
                            emp_data[m_key][day]['pm_in'].append(display_time)
                            
                    elif action == 'C/OUT':
                        if hour < 13:
                            emp_data[m_key][day]['am_out'].append(display_time)
                        else:
                            if not emp_data[m_key][day]['am_out'] and hour < 14:
                                emp_data[m_key][day]['am_out'].append(display_time)
                            else:
                                emp_data[m_key][day]['pm_out'].append(display_time)

        for m_key in emp_data:
            if 'month_name' in m_key or 'year' in m_key:
                continue
            for day in emp_data[m_key]:
                if not day.isdigit():
                    continue
                    
                day_dict = emp_data[m_key][day]
                for field in ['am_in', 'am_out', 'pm_in', 'pm_out']:
                    vals = day_dict[field]
                    
                    unique_vals = []
                    for v in vals:
                        if v not in unique_vals:
                            unique_vals.append(v)
                            
                    if len(unique_vals) == 0:
                        day_dict[field] = ""
                    elif len(unique_vals) == 1:
                        day_dict[field] = unique_vals[0]
                    else:
                        day_dict[field] = unique_vals

    for line in lines:
        line = line.strip()
        if not line:
            continue

        try:
            match = re.match(r'^(.+?)\((\d+)\)$', line)
            if match:
                process_employee_data(current_employee, employee_data, day_checkins)

                if current_employee and employee_data:
                    employees[current_employee] = employee_data

                current_employee = match.group(1).strip()
                employee_data = {}
                day_checkins = {}
                continue

            if not current_employee:
                continue

            parts = line.split()
            if len(parts) < 3:
                continue

            if '/' in parts[0] and parts[0].count('/') == 2:
                date_str = parts[0]
                time_part = parts[1] 
                
                rest_of_line = " ".join(parts[2:]).lower()
                is_pm = 'pm' in rest_of_line
                is_am = 'am' in rest_of_line
                
                action = 'UNKNOWN'
                if 'c/in' in rest_of_line:
                    action = 'C/IN'
                elif 'c/out' in rest_of_line:
                    action = 'C/OUT'
                
                time_part_24 = to_24h(time_part, is_pm, is_am)

                date_parts = date_str.split('/')
                if date_format == "MDY":
                    month_str = date_parts[0]
                    day_str = date_parts[1]
                    year_str = date_parts[2]
                else:
                    day_str = date_parts[0]
                    month_str = date_parts[1]
                    year_str = date_parts[2]

                m_name = calendar.month_name[int(month_str)]
                m_key = f"{m_name} {year_str}"

                if m_key not in employee_data:
                    employee_data[m_key] = {
                        'month_name': m_name,
                        'year': year_str
                    }
                    day_checkins[m_key] = {}

                day_num = int(day_str)
                day_key = str(day_num)

                if day_key not in employee_data[m_key]:
                    employee_data[m_key][day_key] = {
                        'am_in': [], 'am_out': [], 'pm_in': [], 'pm_out': [],
                        'undertime_hrs': '', 'undertime_min': '', 'remarks': '', 'claim_remark': '',
                        'merges': []
                    }
                    day_checkins[m_key][day_key] = []

                if action != 'UNKNOWN':
                    day_checkins[m_key][day_key].append((action, time_part_24))

        except Exception:
            continue

    process_employee_data(current_employee, employee_data, day_checkins)

    if current_employee and employee_data:
        employees[current_employee] = employee_data

    return employees

def generate_dtr(employee_name, employee_data, template_path, approver_name="", approver_title="", period_text="", period_format="full", dtr_format="standard"):
    raw_data = dict(employee_data)
    month_name = raw_data.pop('month_name', '')
    year = raw_data.pop('year', '')

    employee_data = {
        int(k): v for k, v in raw_data.items() if k.isdigit()
    }

    wb = load_workbook(template_path)
    ws = wb.active

    if dtr_format == 'organic':
        name_top = ('C6', 'M6')
        name_bot = ('C55', 'M55')
        approver_cell = ('C61', 'M61')
        approver_title_cell = ('C62', 'M62')
        header_cell = ('E8', 'O8')
        
        left_cols = ['C', 'D', 'E', 'F']
        right_cols = ['M', 'N', 'O', 'P']
        
        left_ut_h, left_ut_m = 'G', 'H'
        right_ut_h, right_ut_m = 'Q', 'R'
        
        left_claim = 'I'
        right_claim = 'S'
        
        left_merge_start = 3 
        left_merge_end = 6 
        right_merge_start = 13 
        right_merge_end = 16 
        
        right_day_col = 'L' 
    else:
        name_top = ('C6', 'K6')
        name_bot = ('C55', 'K55')
        approver_cell = ('C61', 'K61')
        approver_title_cell = ('C62', 'K62')
        header_cell = ('E8', 'M8')
        
        left_cols = ['C', 'D', 'E', 'F']
        right_cols = ['K', 'L', 'M', 'N']
        
        left_ut_h, left_ut_m = 'G', 'H'
        right_ut_h, right_ut_m = 'O', 'P'
        
        left_claim = None
        right_claim = None
        
        left_merge_start = 3 
        left_merge_end = 6 
        right_merge_start = 11 
        right_merge_end = 14 
        
        right_day_col = 'J'

    def safe_write(coord, val, align=None):
        try:
            col_str, row_str = coordinate_from_string(coord)
            col_idx = column_index_from_string(col_str)
            row_idx = int(row_str)
            
            target = ws.cell(row=row_idx, column=col_idx)
            
            if type(target).__name__ == 'MergedCell':
                for mr in ws.merged_cells.ranges:
                    if mr.min_col <= col_idx <= mr.max_col and mr.min_row <= row_idx <= mr.max_row:
                        target = ws.cell(row=mr.min_row, column=mr.min_col)
                        break
            
            if type(target).__name__ != 'MergedCell':
                if isinstance(val, (int, float)):
                    target.value = val
                    target.number_format = '0'
                elif isinstance(val, str) and val.isdigit():
                    target.value = int(val)
                    target.number_format = '0'
                else:
                    target.value = val
                    
                if align:
                    target.alignment = align
        except Exception:
            pass

    safe_write(name_top[0], employee_name)
    safe_write(name_top[1], employee_name)
    safe_write(name_bot[0], employee_name)
    safe_write(name_bot[1], employee_name)

    header_val = period_text if period_text else (f"{month_name} {year}" if month_name else "")
    
    if header_val:
        safe_write(header_cell[0], header_val, Alignment(horizontal='center', vertical='center'))
        safe_write(header_cell[1], header_val, Alignment(horizontal='center', vertical='center'))

    if approver_name:
        safe_write(approver_cell[0], approver_name, Alignment(horizontal='center', vertical='bottom'))
        safe_write(approver_cell[1], approver_name, Alignment(horizontal='center', vertical='bottom'))

    if approver_title:
        safe_write(approver_title_cell[0], approver_title, Alignment(horizontal='center', vertical='center'))
        safe_write(approver_title_cell[1], approver_title, Alignment(horizontal='center', vertical='center'))

    center = Alignment(horizontal='center', vertical='center')
    bold_font = Font(bold=True)
    
    remarks_map = {d: employee_data.get(d, {}).get('remarks', '').strip() for d in range(1, 32)}
    processed_remarks_days = set()

    active_start = 1
    active_end = 31

    if period_format == "1-15":
        active_end = 15
    elif period_format == "16-end":
        active_start = 16

    def resolve_val(v, is_in):
        if isinstance(v, list):
            if len(v) > 0:
                return v[0] if is_in else v[-1]
            return ""
        return v

    for day in range(1, 32):
        row = 13 + day
        if row > 44: break
        
        safe_write(f'B{row}', day)
        safe_write(f'{right_day_col}{row}', day)

        if day < active_start or day > active_end:
            if dtr_format == 'organic':
                cols_to_clear = ['C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T']
            else:
                cols_to_clear = ['C', 'D', 'E', 'F', 'G', 'H', 'K', 'L', 'M', 'N', 'O', 'P']
            for col in cols_to_clear:
                safe_write(f'{col}{row}', "")
            continue

        day_data = employee_data.get(day, {})
        current_remark = remarks_map.get(day, "")

        if day in processed_remarks_days:
            continue

        if current_remark:
            span = 1
            for lookahead in range(day + 1, active_end + 1):
                if remarks_map.get(lookahead) == current_remark:
                    span += 1
                else:
                    break
            
            for d in range(day, day + span):
                processed_remarks_days.add(d)
            
            end_row = row + span - 1
            
            ws.merge_cells(start_row=row, start_column=left_merge_start, end_row=end_row, end_column=left_merge_end)
            target_left = ws.cell(row=row, column=left_merge_start)
            target_left.value = current_remark
            target_left.alignment = center
            target_left.font = bold_font 

            ws.merge_cells(start_row=row, start_column=right_merge_start, end_row=end_row, end_column=right_merge_end)
            target_right = ws.cell(row=row, column=right_merge_start)
            target_right.value = current_remark
            target_right.alignment = center
            target_right.font = bold_font

            for r in range(row, end_row + 1):
                safe_write(f'{left_ut_h}{r}', "")
                safe_write(f'{left_ut_m}{r}', "")
                safe_write(f'{right_ut_h}{r}', "")
                safe_write(f'{right_ut_m}{r}', "")
                if left_claim:
                    safe_write(f'{left_claim}{r}', "")
                    safe_write(f'{right_claim}{r}', "")

        elif day_data.get('merges') and len(day_data.get('merges')) > 0:
            merges = day_data.get('merges', [])
            merged_cols = set()
            
            for m in merges:
                s_col = m['start']
                e_col = m['end']
                m_val = day_data.get(f'merged_{s_col}_{e_col}', '')
                
                ws.merge_cells(start_row=row, start_column=left_merge_start+s_col, end_row=row, end_column=left_merge_start+e_col)
                cell_left = ws.cell(row=row, column=left_merge_start+s_col)
                cell_left.value = m_val
                cell_left.alignment = center
                cell_left.font = bold_font
                
                ws.merge_cells(start_row=row, start_column=right_merge_start+s_col, end_row=row, end_column=right_merge_start+e_col)
                cell_right = ws.cell(row=row, column=right_merge_start+s_col)
                cell_right.value = m_val
                cell_right.alignment = center
                cell_right.font = bold_font
                
                for c in range(s_col, e_col + 1):
                    merged_cols.add(c)
                    
            fields = ['am_in', 'am_out', 'pm_in', 'pm_out']
            
            for i in range(4):
                if i not in merged_cols:
                    val = resolve_val(day_data.get(fields[i], ''), is_in=(i % 2 == 0))
                    safe_write(f'{left_cols[i]}{row}', val, center)
                    safe_write(f'{right_cols[i]}{row}', val, center)
                    
            safe_write(f'{left_ut_h}{row}', day_data.get('undertime_hrs', ''), center)
            safe_write(f'{left_ut_m}{row}', day_data.get('undertime_min', ''), center)
            safe_write(f'{right_ut_h}{row}', day_data.get('undertime_hrs', ''), center)
            safe_write(f'{right_ut_m}{row}', day_data.get('undertime_min', ''), center)
            
            if left_claim:
                safe_write(f'{left_claim}{row}', day_data.get('claim_remark', ''), center)
                safe_write(f'{right_claim}{row}', day_data.get('claim_remark', ''), center)

        else:
            updates = {
                left_cols[0]: resolve_val(day_data.get('am_in', ''), is_in=True),
                left_cols[1]: resolve_val(day_data.get('am_out', ''), is_in=False),
                left_cols[2]: resolve_val(day_data.get('pm_in', ''), is_in=True),
                left_cols[3]: resolve_val(day_data.get('pm_out', ''), is_in=False),
                left_ut_h: day_data.get('undertime_hrs', ''),
                left_ut_m: day_data.get('undertime_min', ''),
                
                right_cols[0]: resolve_val(day_data.get('am_in', ''), is_in=True),
                right_cols[1]: resolve_val(day_data.get('am_out', ''), is_in=False),
                right_cols[2]: resolve_val(day_data.get('pm_in', ''), is_in=True),
                right_cols[3]: resolve_val(day_data.get('pm_out', ''), is_in=False),
                right_ut_h: day_data.get('undertime_hrs', ''),
                right_ut_m: day_data.get('undertime_min', '')
            }
            
            if left_claim:
                updates[left_claim] = day_data.get('claim_remark', '')
                updates[right_claim] = day_data.get('claim_remark', '')

            for col, value in updates.items():
                safe_write(f'{col}{row}', value, center)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generate_ar_from_employee_data(
    employee_name,
    employee_data,
    employee_position,
    employee_office,
    project,
    output_path,
    template_path,
    overrides=None,
    period_text=""
):
    employee_info = {
        "name": employee_name,
        "position": employee_position,
        "office": employee_office,
        "project": project
    }
    
    current_overrides = overrides or {}
    if period_text:
        current_overrides["period_text"] = period_text

    return generate_ar_docx(
        template_path=template_path,
        output_path=output_path,
        employee_data=employee_info,
        parsed_attendance=employee_data,
        overrides=current_overrides
    )

def generate_dtr_adjustment_slip(employee_data, output_path, template_path, overrides={}):
    doc = Document(template_path)

    def append_to_cell(cell, text, newline=False, is_bold=False):
        if not text: return
        text = str(text)
        if not cell.paragraphs:
            p = cell.add_paragraph()
        else:
            p = cell.paragraphs[-1]
        
        if newline:
            run = p.add_run(f"\n{text}")
        else:
            run = p.add_run(f" {text}")
        run.bold = is_bold

    target_table = None
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if "DETAILS OF DTR ADJUSTMENT" in cell.text.upper():
                    target_table = t
                    break
            if target_table: break
        if target_table: break
    
    if not target_table and len(doc.tables) > 0:
        target_table = doc.tables[0]

    if target_table:
        header_flags = {
            'employee_no': False,
            'control_no': False,
            'name': False,
            'filing_date': False
        }

        data_start_idx = 5
        reason_row_idx = 6

        for i, row in enumerate(target_table.rows):
            row_text = " ".join([c.text.upper() for c in row.cells])
            if "DETAILS OF DTR ADJUSTMENT" in row_text:
                data_start_idx = i + 3 
                break
        
        for i in range(data_start_idx, len(target_table.rows)):
            row_text = " ".join([c.text.upper() for c in target_table.rows[i].cells])
            if "REASON" in row_text:
                reason_row_idx = i
                break

        for i in range(data_start_idx):
            distinct_cells = []
            for c in target_table.rows[i].cells:
                if c not in distinct_cells:
                    distinct_cells.append(c)
            
            for cell in distinct_cells:
                txt = cell.text.upper()
                if "EMPLOYEE NO" in txt and not header_flags['employee_no']:
                    append_to_cell(cell, overrides.get('employee_no', ''), newline=False, is_bold=False)
                    header_flags['employee_no'] = True
                elif "CONTROL NO" in txt and not header_flags['control_no']:
                    append_to_cell(cell, overrides.get('control_no', ''), newline=False, is_bold=False)
                    header_flags['control_no'] = True
                elif "NAME" in txt and "SIGNATURE" not in txt and not header_flags['name'] and "PRINTED" not in txt:
                    append_to_cell(cell, overrides.get('adjustment_name', '').upper(), newline=False, is_bold=True)
                    header_flags['name'] = True
                elif "DATE/TIME" in txt and not header_flags['filing_date']:
                    fdate = overrides.get('filing_date', '')
                    if fdate:
                        try:
                            dt = datetime.strptime(fdate, "%Y-%m-%dT%H:%M")
                            fdate = dt.strftime("%B %d, %Y %I:%M %p").upper()
                        except ValueError:
                            fdate = fdate.upper()
                    append_to_cell(cell, fdate, newline=False, is_bold=True)
                    header_flags['filing_date'] = True

        adjustment_rows = overrides.get('adjustment_rows', [])
        valid_rows = [r for r in adjustment_rows if any(str(r.get(k, '')).strip() for k in ['date', 'am_in', 'am_out', 'pm_in', 'pm_out', 'evening_in', 'evening_out'])]
        
        current_row_idx = data_start_idx

        for i, row_data in enumerate(valid_rows):
            if current_row_idx >= reason_row_idx:
                ref_row = target_table.rows[current_row_idx - 1]
                new_row_xml = copy.deepcopy(ref_row._element)
                
                for tc in new_row_xml.tc_lst:
                    for p in tc.p_lst:
                        p.clear_content()
                
                ref_row._element.addnext(new_row_xml)
                reason_row_idx += 1 
            
            row = target_table.rows[current_row_idx]
            
            distinct_cells = []
            for c in row.cells:
                if c not in distinct_cells:
                    distinct_cells.append(c)
                    
            dc_count = len(distinct_cells)
            
            def set_small_text(cell, txt):
                cell.text = ''
                p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
                run = p.add_run(str(txt) if txt else '')
                run.font.size = Pt(8) 
                p.alignment = 1 

            if dc_count > 0: set_small_text(distinct_cells[0], row_data.get('date', ''))
            if dc_count > 1: set_small_text(distinct_cells[1], to_12h(row_data.get('am_in', '')))
            if dc_count > 2: set_small_text(distinct_cells[2], to_12h(row_data.get('am_out', '')))
            if dc_count > 3: set_small_text(distinct_cells[3], to_12h(row_data.get('pm_in', '')))
            if dc_count > 4: set_small_text(distinct_cells[4], to_12h(row_data.get('pm_out', '')))
            if dc_count > 5: set_small_text(distinct_cells[5], to_12h(row_data.get('evening_in', '')))
            if dc_count > 6: set_small_text(distinct_cells[6], to_12h(row_data.get('evening_out', '')))
            
            current_row_idx += 1

        if reason_row_idx < len(target_table.rows):
            reason_key = overrides.get('reason', '')
            ob_with = overrides.get('ob_with', '')
            ob_at = overrides.get('ob_at', '')
            personal_details = overrides.get('personal_details', '')
            other_details = overrides.get('other_details', '')

            target_row = target_table.rows[reason_row_idx]
            target_cell = target_row.cells[0]
            for cell in target_row.cells:
                if "REASON" in cell.text.upper():
                    target_cell = cell
                    break
            
            inner_table = target_cell.add_table(rows=2, cols=2)
            
            def add_checkbox(paragraph, is_checked, text):
                box_char = "■" if is_checked else "□"
                run_box = paragraph.add_run(box_char)
                run_box.font.name = 'Arial'
                run_box._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
                run_box._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
                paragraph.add_run(f" {text}")

            c00 = inner_table.cell(0, 0)
            p00 = c00.paragraphs[0] if c00.paragraphs else c00.add_paragraph()
            add_checkbox(p00, reason_key == 'fingerprint', "Fingerprint not recognized")
            
            c01 = inner_table.cell(0, 1)
            p01 = c01.paragraphs[0] if c01.paragraphs else c01.add_paragraph()
            add_checkbox(p01, reason_key == 'personal', "Personal Reason:")
            if reason_key == 'personal':
                r_p = p01.add_run(f" {personal_details}" if personal_details else "                          ")
                r_p.underline = True
            
            c10 = inner_table.cell(1, 0)
            p10 = c10.paragraphs[0] if c10.paragraphs else c10.add_paragraph()
            add_checkbox(p10, reason_key == 'ob', "On Official Business/Pass Slip")
            
            p10_where = c10.add_paragraph()
            p10_where.add_run("    With: ")
            if reason_key == 'ob':
                r_w = p10_where.add_run(ob_with if ob_with else "                          ")
                r_w.underline = True
                
            p10_at = c10.add_paragraph()
            p10_at.add_run("    At: ")
            if reason_key == 'ob':
                r_a = p10_at.add_run(ob_at if ob_at else "                          ")
                r_a.underline = True
            
            c11 = inner_table.cell(1, 1)
            p11 = c11.paragraphs[0] if c11.paragraphs else c11.add_paragraph()
            add_checkbox(p11, reason_key == 'other', "Other reasons:")
            if reason_key == 'other':
                r_ot = p11.add_run(f" {other_details}" if other_details else "                          ")
                r_ot.underline = True

        cert_row_idx = -1
        for i in range(reason_row_idx, len(target_table.rows)):
            row_text = " ".join([c.text.upper() for c in target_table.rows[i].cells])
            if "CERTIFIED TRUE AND CORRECT" in row_text:
                cert_row_idx = i
                break
        
        if cert_row_idx != -1 and cert_row_idx + 1 < len(target_table.rows):
            name_row = target_table.rows[cert_row_idx + 1]
            emp_name = overrides.get('name', '').upper()
            approver = overrides.get('approver', '').upper()
            
            distinct_cells = []
            for c in name_row.cells:
                if c not in distinct_cells:
                    distinct_cells.append(c)
            
            if len(distinct_cells) >= 2:
                distinct_cells[0].text = ''
                p0 = distinct_cells[0].paragraphs[0] if distinct_cells[0].paragraphs else distinct_cells[0].add_paragraph()
                run0 = p0.add_run(emp_name)
                run0.bold = True
                p0.alignment = 1

                distinct_cells[1].text = ''
                p1 = distinct_cells[1].paragraphs[0] if distinct_cells[1].paragraphs else distinct_cells[1].add_paragraph()
                run1 = p1.add_run(approver)
                run1.bold = True
                p1.alignment = 1

    doc.save(output_path)
    return output_path